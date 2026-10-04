"""SQLite initialization for the Python API."""

import os
import sqlite3
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parent.parent
_configured_db_path = Path(
    os.getenv("DB_PATH", str(PROJECT_ROOT / "server" / "db" / "shopsense.db"))
)
DB_PATH = (
    _configured_db_path
    if _configured_db_path.is_absolute()
    else PROJECT_ROOT / _configured_db_path
).resolve()

SCHEMA = """
CREATE TABLE IF NOT EXISTS admins (
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  name TEXT NOT NULL,
  email TEXT UNIQUE NOT NULL,
  password TEXT NOT NULL,
  email_verified INTEGER NOT NULL DEFAULT 0
);

CREATE TABLE IF NOT EXISTS vendors (
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  full_name TEXT NOT NULL,
  business_name TEXT NOT NULL,
  email TEXT UNIQUE NOT NULL,
  password TEXT NOT NULL,
  phone TEXT,
  business_address TEXT,
  status TEXT NOT NULL DEFAULT 'pending',
  joined_at TEXT NOT NULL DEFAULT (datetime('now')),
  email_verified INTEGER NOT NULL DEFAULT 0
);

CREATE TABLE IF NOT EXISTS products (
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  vendor_id INTEGER NOT NULL,
  name TEXT NOT NULL,
  description TEXT,
  category TEXT NOT NULL,
  price REAL NOT NULL,
  stock INTEGER NOT NULL DEFAULT 0,
  image_url TEXT,
  created_at TEXT NOT NULL DEFAULT (datetime('now')),
  seo_title TEXT NOT NULL DEFAULT '',
  short_description TEXT NOT NULL DEFAULT '',
  meta_title TEXT NOT NULL DEFAULT '',
  meta_description TEXT NOT NULL DEFAULT '',
  seo_keywords TEXT NOT NULL DEFAULT '',
  product_tags TEXT NOT NULL DEFAULT '',
  key_features TEXT NOT NULL DEFAULT '',
  vendor_hints TEXT NOT NULL DEFAULT '',
  FOREIGN KEY (vendor_id) REFERENCES vendors(id) ON DELETE CASCADE
);

CREATE TABLE IF NOT EXISTS sales (
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  vendor_id INTEGER NOT NULL,
  product_id INTEGER NOT NULL,
  quantity INTEGER NOT NULL,
  amount REAL NOT NULL,
  sold_at TEXT NOT NULL DEFAULT (datetime('now')),
  FOREIGN KEY (vendor_id) REFERENCES vendors(id) ON DELETE CASCADE
);

CREATE TABLE IF NOT EXISTS app_settings (
  setting_key TEXT PRIMARY KEY,
  setting_value TEXT NOT NULL
);

CREATE TABLE IF NOT EXISTS analytics_products (
  product_id TEXT PRIMARY KEY,
  product_name TEXT NOT NULL,
  category TEXT NOT NULL,
  price REAL NOT NULL,
  stock INTEGER NOT NULL
);

CREATE TABLE IF NOT EXISTS analytics_customers (
  customer_id TEXT PRIMARY KEY,
  customer_name TEXT NOT NULL,
  email TEXT NOT NULL UNIQUE
);

CREATE TABLE IF NOT EXISTS analytics_orders (
  order_id TEXT PRIMARY KEY,
  customer_id TEXT NOT NULL,
  order_date TEXT NOT NULL,
  total_amount REAL NOT NULL,
  FOREIGN KEY (customer_id) REFERENCES analytics_customers(customer_id)
);

CREATE TABLE IF NOT EXISTS analytics_order_items (
  id INTEGER PRIMARY KEY AUTOINCREMENT,
  order_id TEXT NOT NULL,
  product_id TEXT NOT NULL,
  quantity INTEGER NOT NULL,
  unit_price REAL NOT NULL,
  FOREIGN KEY (order_id) REFERENCES analytics_orders(order_id),
  FOREIGN KEY (product_id) REFERENCES analytics_products(product_id)
);

CREATE TABLE IF NOT EXISTS email_verifications (
  account_type TEXT NOT NULL CHECK (account_type IN ('admin', 'vendor')),
  account_id INTEGER NOT NULL,
  token_hash TEXT NOT NULL UNIQUE,
  expires_at TEXT NOT NULL,
  sent_at TEXT NOT NULL,
  verified_at TEXT,
  PRIMARY KEY (account_type, account_id)
);

CREATE INDEX IF NOT EXISTS idx_analytics_orders_customer ON analytics_orders(customer_id);
CREATE INDEX IF NOT EXISTS idx_analytics_orders_date ON analytics_orders(order_date);
CREATE INDEX IF NOT EXISTS idx_analytics_order_items_product ON analytics_order_items(product_id);
CREATE INDEX IF NOT EXISTS idx_analytics_order_items_order ON analytics_order_items(order_id);
CREATE INDEX IF NOT EXISTS idx_products_vendor_stock ON products(vendor_id, stock, name);
CREATE INDEX IF NOT EXISTS idx_sales_vendor_date ON sales(vendor_id, sold_at);
CREATE INDEX IF NOT EXISTS idx_sales_vendor_product ON sales(vendor_id, product_id);
"""


def _columns(connection: sqlite3.Connection, table: str) -> set[str]:
    return {row["name"] for row in connection.execute(f"PRAGMA table_info({table})")}


def _import_analytics_dataset(connection: sqlite3.Connection) -> None:
    if connection.execute("SELECT 1 FROM analytics_orders LIMIT 1").fetchone():
        return

    dataset_dir = PROJECT_ROOT / "dataset"
    table_specs = (
        ("products_1.xlsx", "analytics_products", ("product_id", "product_name", "category", "price", "stock")),
        ("customers_1.xlsx", "analytics_customers", ("customer_id", "customer_name", "email")),
        ("orders_1.xlsx", "analytics_orders", ("order_id", "customer_id", "order_date", "total_amount")),
        ("order_items_1.xlsx", "analytics_order_items", ("order_id", "product_id", "quantity", "unit_price")),
    )
    if not all((dataset_dir / filename).is_file() for filename, _, _ in table_specs):
        return

    with connection:
        for filename, table, columns in table_specs:
            workbook = load_workbook(dataset_dir / filename, read_only=True, data_only=True)
            try:
                sheet = workbook[workbook.sheetnames[0]]
                rows = sheet.iter_rows(values_only=True)
                headers = [str(value).strip() if value is not None else "" for value in next(rows)]
                missing = set(columns) - set(headers)
                if missing:
                    raise ValueError(f"{filename} is missing required columns: {', '.join(sorted(missing))}")
                indexes = [headers.index(column) for column in columns]
                placeholders = ", ".join("?" for _ in columns)
                quoted_columns = ", ".join(f'"{column}"' for column in columns)
                insert = f'INSERT OR IGNORE INTO "{table}" ({quoted_columns}) VALUES ({placeholders})'
                for row in rows:
                    values: list[Any] = [row[index] if index < len(row) else None for index in indexes]
                    if table in {"analytics_products", "analytics_orders", "analytics_order_items"}:
                        numeric_columns = {
                            "analytics_products": {"price", "stock"},
                            "analytics_orders": {"total_amount"},
                            "analytics_order_items": {"quantity", "unit_price"},
                        }[table]
                        values = [
                            float(value) if column in numeric_columns and value is not None else value
                            for column, value in zip(columns, values)
                        ]
                    connection.execute(insert, values)
            finally:
                workbook.close()


def initialize_database() -> None:
    """Create application tables and migrate databases created by the old API."""
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(DB_PATH)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    connection.execute("PRAGMA journal_mode = WAL")
    try:
        connection.executescript(SCHEMA)
        for table in ("admins", "vendors"):
            if "email_verified" not in _columns(connection, table):
                connection.execute(
                    f"ALTER TABLE {table} ADD COLUMN email_verified INTEGER NOT NULL DEFAULT 0"
                )
        product_migrations = {
            "image_url": "TEXT",
            "seo_title": "TEXT NOT NULL DEFAULT ''",
            "short_description": "TEXT NOT NULL DEFAULT ''",
            "meta_title": "TEXT NOT NULL DEFAULT ''",
            "meta_description": "TEXT NOT NULL DEFAULT ''",
            "seo_keywords": "TEXT NOT NULL DEFAULT ''",
            "product_tags": "TEXT NOT NULL DEFAULT ''",
            "key_features": "TEXT NOT NULL DEFAULT ''",
            "vendor_hints": "TEXT NOT NULL DEFAULT ''",
        }
        product_columns = _columns(connection, "products")
        for column, definition in product_migrations.items():
            if column not in product_columns:
                connection.execute(f"ALTER TABLE products ADD COLUMN {column} {definition}")
        connection.commit()
        _import_analytics_dataset(connection)
    finally:
        connection.close()
